"""Genera plantilla/certificado_base.png a partir de la diapositiva 2 del pptx.

Quita el cuadro de texto central (CERTIFICAMOS QUE ... fecha), deja fondo,
medallón, logo y firmas con sus textos de pie, y renderiza con LibreOffice
headless. También extrae el logo de Towers para el sitio público.

Uso:  python plantilla/crear_base.py [--original]
  --original  además renderiza la diapositiva 2 SIN modificar (referencia para calibrar).
"""
import io
import shutil
import subprocess
import sys
from pathlib import Path

import pymupdf
from PIL import Image
from pptx import Presentation

RAIZ = Path(__file__).resolve().parent.parent
PPTX = RAIZ / "Ejemplos de certificados.pptx"
SALIDA = RAIZ / "plantilla"
ANCHO_PX = 3200
SOFFICE_CANDIDATOS = [
    shutil.which("soffice"),
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
]


def soffice():
    for c in SOFFICE_CANDIDATOS:
        if c and Path(c).exists():
            return c
    return None


def solo_diapositiva_2(prs):
    """Elimina todas las diapositivas salvo la 2."""
    ids = prs.slides._sldIdLst
    for i, sld in reversed(list(enumerate(list(ids)))):
        if i != 1:
            prs.part.drop_rel(sld.rId)
            ids.remove(sld)


def render_powerpoint(pptx_path: Path, png_path: Path):
    """Alternativa cuando LibreOffice no está disponible: exporta con PowerPoint (COM)."""
    ps = (
        "$p = New-Object -ComObject PowerPoint.Application;"
        f"$d = $p.Presentations.Open('{pptx_path}', $true, $false, $false);"
        f"$d.Slides.Item(1).Export('{png_path}', 'PNG', {ANCHO_PX}, {ANCHO_PX * 9 // 16});"
        "$d.Close(); $p.Quit()"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True, timeout=300)
    print(f"{png_path.name}: {Image.open(png_path).size} (PowerPoint)")


def render(pptx_path: Path, png_path: Path):
    if soffice() is None:
        return render_powerpoint(pptx_path, png_path)
    tmp = SALIDA / "_tmp"
    tmp.mkdir(exist_ok=True)
    subprocess.run(
        [soffice(), "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(pptx_path)],
        check=True, capture_output=True, timeout=300,
    )
    pdf = tmp / (pptx_path.stem + ".pdf")
    doc = pymupdf.open(pdf)
    page = doc[0]
    zoom = ANCHO_PX / page.rect.width
    pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
    pix.save(png_path)
    print(f"{png_path.name}: {pix.width}x{pix.height}")
    doc.close()


def extraer_logo(prs):
    """El logo de Towers es la imagen cuadrada grande de la diapositiva 2."""
    for sh in prs.slides[1].shapes:
        if sh.shape_type == 13 and sh.image.size == (2400, 2400):
            img = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA")
            # el logo es un PNG con fondo transparente: se recorta por el canal alfa
            x0, y0, x1, y1 = img.getchannel("A").point(lambda v: 255 if v > 40 else 0).getbbox()
            pad = 40
            img = img.crop((max(x0 - pad, 0), max(y0 - pad, 0), min(x1 + pad, 2400), min(y1 + pad, 2400)))
            img.thumbnail((700, 700), Image.LANCZOS)
            destino = RAIZ / "publico" / "assets" / "logo.png"
            destino.parent.mkdir(parents=True, exist_ok=True)
            img.save(destino, optimize=True)
            print(f"logo: {img.size} -> {destino}")
            return
    sys.exit("No se encontró la imagen del logo.")


def main():
    if "--original" in sys.argv:
        prs = Presentation(PPTX)
        solo_diapositiva_2(prs)
        p = SALIDA / "_original_diap2.pptx"
        prs.save(p)
        render(p, SALIDA / "_original_diap2.png")

    prs = Presentation(PPTX)
    extraer_logo(prs)
    solo_diapositiva_2(prs)
    diap = prs.slides[0]
    quitados = 0
    for sh in list(diap.shapes):
        if sh.has_text_frame and "CERTIFICAMOS QUE" in sh.text_frame.text:
            sh._element.getparent().remove(sh._element)
            quitados += 1
    assert quitados == 1, "no se encontró el cuadro de texto central"
    base = SALIDA / "base_sin_texto.pptx"
    prs.save(base)
    render(base, SALIDA / "certificado_base.png")
    shutil.rmtree(SALIDA / "_tmp", ignore_errors=True)


if __name__ == "__main__":
    main()
